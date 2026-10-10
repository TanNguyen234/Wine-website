import requests
import io
import openpyxl
from pathlib import Path
from PIL import Image

def run_excel_checks():
    print("=== CHECK 1: EXCEL OPENXML INTEGRITY ===")
    s = requests.Session()
    s.get("http://localhost:8080/login")
    login_resp = s.post("http://localhost:8080/login", data={"username": "admin", "password": "admin123"}, allow_redirects=True)
    assert login_resp.status_code == 200

    # 1. Template
    r_tpl = s.get("http://localhost:8080/api/admin/excel/template")
    print(f"Template Status: {r_tpl.status_code}")
    print(f"Template Content-Type: {r_tpl.headers.get('content-type')}")
    print(f"Template Byte Length: {len(r_tpl.content)}")
    assert r_tpl.content[:4] == b"PK\x03\x04", "Header must be PK\x03\x04"
    wb_tpl = openpyxl.load_workbook(io.BytesIO(r_tpl.content))
    print(f"Template Sheets: {wb_tpl.sheetnames}")
    rows_tpl = list(wb_tpl.active.iter_rows(values_only=True))
    print(f"Template Rows: {len(rows_tpl)}, Header: {rows_tpl[0]}")

    # 2. Inventory Export
    r_inv = s.get("http://localhost:8080/api/admin/excel/export/inventory")
    print(f"Inventory Status: {r_inv.status_code}")
    print(f"Inventory Content-Type: {r_inv.headers.get('content-type')}")
    print(f"Inventory Byte Length: {len(r_inv.content)}")
    assert r_inv.content[:4] == b"PK\x03\x04", "Header must be PK\x03\x04"
    wb_inv = openpyxl.load_workbook(io.BytesIO(r_inv.content))
    print(f"Inventory Sheets: {wb_inv.sheetnames}")
    rows_inv = list(wb_inv.active.iter_rows(values_only=True))
    print(f"Inventory Rows: {len(rows_inv)}, Header: {rows_inv[0]}")

def run_screenshot_checks():
    print("\n=== CHECK 2: SCREENSHOT FORENSICS ===")
    ss_dir = Path(r"D:\Projects\strongwine_2\Wine-website\test-results\screenshots")
    pngs = sorted(list(ss_dir.glob("*.png")))
    print(f"Found {len(pngs)} screenshot files.")
    assert len(pngs) == 23, f"Expected 23 screenshots, got {len(pngs)}"
    for p in pngs:
        with open(p, "rb") as f:
            sig = f.read(8)
        assert sig == b"\x89PNG\r\n\x1a\n", f"Invalid PNG magic bytes in {p.name}"
        im = Image.open(p)
        print(f"  {p.name}: {p.stat().st_size} bytes | {im.size} | {im.mode}")

if __name__ == "__main__":
    run_excel_checks()
    run_screenshot_checks()
    print("\nALL FORENSIC CHECKS PASSED EMPIRICALLY!")
